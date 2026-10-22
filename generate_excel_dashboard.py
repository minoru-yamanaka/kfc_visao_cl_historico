import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import pandas as pd
import json

with open('data.json', 'r', encoding='utf-8') as f:
    stores = json.load(f)

wb = openpyxl.Workbook()

# ----------------- ABA 1: ROADMAP GANTT & DIAS ACUMULADOS -----------------
ws = wb.active
ws.title = "Roadmap Expansão 2026"
ws.views.sheetView[0].showGridLines = True

# Paleta HEX (sem # para openpyxl)
HEX_PURPLE = "8B5CF6"
HEX_GREEN = "16A34A"
HEX_RED = "EF4444"
HEX_ORANGE = "FB923C"
HEX_YELLOW = "FACC15"
HEX_DARK_GREEN = "15803D"
HEX_GRAY = "94A3B8"
HEX_KFC_RED = "E4002B"
HEX_DARK_BG = "0F172A"

# Estilos
font_title = Font(name="Calibri", size=16, bold=True, color="FFFFFF")
font_sub = Font(name="Calibri", size=11, italic=True, color="E2E8F0")
font_header = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
font_cell = Font(name="Calibri", size=10)
font_cell_bold = Font(name="Calibri", size=10, bold=True)
font_white_bold = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
font_dark_bold = Font(name="Calibri", size=10, bold=True, color="0F172A")

fill_kfc_red = PatternFill(start_color=HEX_KFC_RED, end_color=HEX_KFC_RED, fill_type="solid")
fill_dark = PatternFill(start_color=HEX_DARK_BG, end_color=HEX_DARK_BG, fill_type="solid")
fill_light_gray = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")

thin_side = Side(border_style="thin", color="CBD5E1")
border_cell = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
border_header = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)

# Cabeçalho Principal
ws.merge_cells("A1:N1")
ws["A1"] = "KFC 2026 — ROADMAP DE EXPANSÃO | MONITORAMENTO DE DIAS ACUMULADOS POR SHOPPING"
ws["A1"].font = font_title
ws["A1"].fill = fill_kfc_red
ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
ws.row_dimensions[1].height = 35

ws.merge_cells("A2:N2")
ws["A2"] = "Formatação Condicional Rigorosa por Etapa: Roxo (<0d E1), Verde (<=7d E1 / <=30d E2), Vermelho (Atraso), Laranja Claro (<=30d E3), Amarelo (<90d E4), Verde Escuro (>=90d E4)"
ws["A2"].font = font_sub
ws["A2"].fill = fill_dark
ws["A2"].alignment = Alignment(horizontal="center", vertical="center")
ws.row_dimensions[2].height = 24

# Linha 4: Cabeçalhos das Colunas
headers = [
    "ID", "Shopping / Localidade", "Fmt",
    "Assinatura Contrato", "Solicitação CNPJ", "Emissão CNPJ", "Criação SAP", "Inauguração",
    "E1: Contrato ➔ Solic. CNPJ", "E2: Solic. ➔ Emissão CNPJ", "E3: Emissão ➔ Criação SAP", "E4: SAP ➔ Inauguração",
    "Total Dias Acumulados", "Status Executivo"
]

ws.row_dimensions[4].height = 28
for col_idx, h in enumerate(headers, 1):
    cell = ws.cell(row=4, column=col_idx, value=h)
    cell.font = font_header
    cell.fill = fill_dark
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = border_header

# Preencher dados
for row_idx, s in enumerate(stores, 5):
    ws.row_dimensions[row_idx].height = 22
    
    # ID
    c_id = ws.cell(row=row_idx, column=1, value=s["id"])
    c_id.alignment = Alignment(horizontal="center", vertical="center")
    c_id.border = border_cell
    
    # Shopping
    c_loja = ws.cell(row=row_idx, column=2, value=s["loja"])
    c_loja.font = font_cell_bold
    c_loja.border = border_cell
    
    # Formato
    c_fmt = ws.cell(row=row_idx, column=3, value=s["formato"])
    c_fmt.alignment = Alignment(horizontal="center", vertical="center")
    c_fmt.border = border_cell
    
    # Datas
    dts = [s["datas"]["contrato"], s["datas"]["solicitacao"], s["datas"]["emissao"], s["datas"]["sap"], s["datas"]["inauguracao"]]
    for d_idx, dt in enumerate(dts, 4):
        c_dt = ws.cell(row=row_idx, column=d_idx, value=dt[:10] if dt and dt not in ["Não Possui", "Plano de Ação"] else dt)
        c_dt.alignment = Alignment(horizontal="center", vertical="center")
        c_dt.border = border_cell
    
    # Etapas com coloração condicional rigorosa
    stages_info = [
        (9, s["etapa1"]),
        (10, s["etapa2"]),
        (11, s["etapa3"]),
        (12, s["etapa4"])
    ]
    
    for col_i, st in stages_info:
        val = st["dias"]
        cell_st = ws.cell(row=row_idx, column=col_i, value=f"{val}d" if isinstance(val, (int, float)) else val)
        cell_st.alignment = Alignment(horizontal="center", vertical="center")
        cell_st.border = border_cell
        
        # Cor de fundo exata
        c_class = st["class"]
        if c_class == "purple":
            cell_st.fill = PatternFill(start_color=HEX_PURPLE, end_color=HEX_PURPLE, fill_type="solid")
            cell_st.font = font_white_bold
        elif c_class == "green":
            cell_st.fill = PatternFill(start_color=HEX_GREEN, end_color=HEX_GREEN, fill_type="solid")
            cell_st.font = font_white_bold
        elif c_class == "red":
            cell_st.fill = PatternFill(start_color=HEX_RED, end_color=HEX_RED, fill_type="solid")
            cell_st.font = font_white_bold
        elif c_class == "light-orange":
            cell_st.fill = PatternFill(start_color=HEX_ORANGE, end_color=HEX_ORANGE, fill_type="solid")
            cell_st.font = font_white_bold
        elif c_class == "yellow-alert":
            cell_st.fill = PatternFill(start_color=HEX_YELLOW, end_color=HEX_YELLOW, fill_type="solid")
            cell_st.font = font_dark_bold
        elif c_class == "dark-green":
            cell_st.fill = PatternFill(start_color=HEX_DARK_GREEN, end_color=HEX_DARK_GREEN, fill_type="solid")
            cell_st.font = font_white_bold
        else:
            cell_st.fill = PatternFill(start_color=HEX_GRAY, end_color=HEX_GRAY, fill_type="solid")
            cell_st.font = font_dark_bold

    # Total de Dias Acumulados
    tot_label = f"{s['total_acumulado']}d total" if s["todas_etapas_concluidas"] else f"{s['total_acumulado']}d + Ação"
    c_tot = ws.cell(row=row_idx, column=13, value=tot_label)
    c_tot.font = font_cell_bold
    c_tot.fill = fill_light_gray
    c_tot.alignment = Alignment(horizontal="center", vertical="center")
    c_tot.border = border_cell

    # Status Executivo
    if not s["todas_etapas_concluidas"]:
        status_txt = "📋 Plano de Ação em Aberto"
    elif s["etapa4"]["class"] == "yellow-alert":
        status_txt = "⚠️ Alto Risco Inauguração (< 90d)"
    elif s["etapa1"]["class"] == "purple":
        status_txt = "⚡ Antecipação Estratégica"
    else:
        status_txt = "✅ Cronograma Concluído"
        
    c_status = ws.cell(row=row_idx, column=14, value=status_txt)
    c_status.font = font_cell
    c_status.alignment = Alignment(horizontal="left", vertical="center")
    c_status.border = border_cell

# ----------------- LEGENDA NA PLANILHA -----------------
start_leg = len(stores) + 7
ws.cell(row=start_leg, column=2, value="LEGENDA OFICIAL DE CORES").font = font_header
ws.cell(row=start_leg, column=2).fill = fill_kfc_red
ws.cell(row=start_leg, column=2).alignment = Alignment(horizontal="center")

leg_items = [
    (HEX_PURPLE, "Roxo", "Etapa 1: Dias < 0 (Solicitação feita antes do contrato para ganhar tempo)"),
    (HEX_GREEN, "Verde", "Etapa 1: ≤ 7 dias (No prazo) | Etapa 2: ≤ 30 dias (No prazo)"),
    (HEX_RED, "Vermelho", "Atraso: Etapa 1 (> 7 dias) | Etapa 2 (> 30 dias) | Etapa 3 (> 30 dias)"),
    (HEX_ORANGE, "Laranja Claro", "Etapa 3: ≤ 30 dias (No prazo para criação SAP)"),
    (HEX_YELLOW, "Amarelo / Alerta", "Etapa 4: < 90 dias restantes (Alto risco para a inauguração)"),
    (HEX_DARK_GREEN, "Verde Escuro", "Etapa 4: ≥ 90 dias restantes (Prazo seguro e saudável)"),
    (HEX_GRAY, "Cinza", "Plano de Ação: Etapa pendente de execução / cadastro")
]

for l_idx, (color, name, desc) in enumerate(leg_items, start_leg + 1):
    c_swatch = ws.cell(row=l_idx, column=2, value=name)
    c_swatch.fill = PatternFill(start_color=color, end_color=color, fill_type="solid")
    c_swatch.font = font_white_bold if color != HEX_YELLOW else font_dark_bold
    c_swatch.alignment = Alignment(horizontal="center", vertical="center")
    c_swatch.border = border_cell
    
    ws.merge_cells(start_row=l_idx, start_column=3, end_row=l_idx, end_column=9)
    c_desc = ws.cell(row=l_idx, column=3, value=desc)
    c_desc.font = font_cell
    c_desc.alignment = Alignment(horizontal="left", vertical="center")

# Ajuste automático das larguras das colunas
for col in ws.columns:
    max_len = 0
    col_letter = get_column_letter(col[0].column)
    for cell in col:
        if cell.row > 2 and cell.value:
            max_len = max(max_len, len(str(cell.value)))
    ws.column_dimensions[col_letter].width = max(max_len + 3, 11)

ws.column_dimensions["B"].width = 32
ws.column_dimensions["N"].width = 30

excel_file = "KFC_2026_Roadmap_Dashboard.xlsx"
wb.save(excel_file)
print(f"Planilha {excel_file} gerada com sucesso!")
